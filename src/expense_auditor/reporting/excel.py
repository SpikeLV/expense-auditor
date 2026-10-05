"""Excel workbook for a monthly expense audit.

The workbook is built from a ``ReconciliationResult``. Incoming transactions
are included when the caller supplies the original transaction list, because
the result itself stores outgoing rows only. Transaction has no separate
value date or reference field; the posted date and the description are the
source values.
"""

from __future__ import annotations

import os
from collections.abc import Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import Cell
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from expense_auditor.models.document import Document
from expense_auditor.models.match import Match, MatchStatus
from expense_auditor.models.reconciliation import (
    ReconciliationResult,
    ReconciliationSummary,
)
from expense_auditor.models.transaction import Transaction, TransactionDirection

_AMOUNT_FORMAT = "#,##0.00"
_PERCENT_FORMAT = "0.0%"
_DATE_FORMAT = "YYYY-MM-DD"
_TIMESTAMP_FORMAT = "YYYY-MM-DD HH:MM"
_SCORE_FORMAT = "0.000"
_MAX_COLUMN_WIDTH = 48

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_SECTION_FONT = Font(bold=True, color="1F4E79", size=13)
_TITLE_FONT = Font(bold=True, color="1F4E79", size=16)

_CONFIRMED_FILL = PatternFill("solid", fgColor="C6EFCE")
_PROBABLE_FILL = PatternFill("solid", fgColor="FFEB9C")
_MISSING_FILL = PatternFill("solid", fgColor="FFC7CE")
_CONFIRMED_FONT = Font(color="006100")
_PROBABLE_FONT = Font(color="9C5700")
_MISSING_FONT = Font(color="9C0006")

_TRANSACTION_HEADERS = (
    "Transaction ID",
    "Date",
    "Direction",
    "Amount",
    "Currency",
    "Merchant",
    "Merchant Normalized",
    "Description",
    "Status",
    "Match Score",
    "Matched Documents",
    "Matched Document IDs",
    "Match Reason",
    "Amount Score",
    "Currency Score",
    "Merchant Score",
    "Date Score",
    "Reference Score",
    "Combination Match",
    "Number of Documents",
)
_MISSING_HEADERS = (
    "Transaction ID",
    "Date",
    "Amount",
    "Currency",
    "Merchant",
    "Description",
    "Status",
    "Match Score",
    "Reason",
)
_PROBABLE_HEADERS = (
    "Transaction ID",
    "Date",
    "Amount",
    "Currency",
    "Merchant",
    "Description",
    "Status",
    "Match Score",
    "Matched Documents",
    "Matched Document IDs",
    "Reason",
    "Amount Score",
    "Currency Score",
    "Merchant Score",
    "Date Score",
    "Reference Score",
    "Combination Match",
    "Number of Documents",
)
_UNMATCHED_HEADERS = (
    "Document ID",
    "Filename",
    "Document Type",
    "Document Date",
    "Merchant",
    "Merchant Normalized",
    "Amount",
    "Currency",
    "Invoice Number",
    "Extraction Method",
    "File Hash",
    "File Path",
)
_DOCUMENT_HEADERS = (
    "Document ID",
    "Filename",
    "Document Type",
    "Document Date",
    "Merchant Raw",
    "Merchant Normalized",
    "Amount",
    "Currency",
    "Invoice Number",
    "Extraction Method",
    "File Hash",
    "File Path",
    "Match Usage",
    "Matched Transaction IDs",
)
_DUPLICATE_HEADERS = (
    "Document ID",
    "Filename",
    "Document Date",
    "Amount",
    "Currency",
    "Transaction Count",
    "Matched Transaction IDs",
)

_SCORE_HEADERS = frozenset(
    {
        "Match Score",
        "Amount Score",
        "Currency Score",
        "Merchant Score",
        "Date Score",
        "Reference Score",
    }
)
_DATE_HEADERS = frozenset({"Date", "Document Date"})
_AMOUNT_HEADERS = frozenset({"Amount"})

_LEGEND = (
    (
        MatchStatus.CONFIRMED.value,
        "Strong supporting evidence: exact amount and currency, plus a strong "
        "merchant or an exact reference.",
    ),
    (
        MatchStatus.PROBABLE.value,
        "Plausible match requiring review. Credible evidence is present, and "
        "the confirmed rule is not met.",
    ),
    (
        MatchStatus.MISSING.value,
        "No sufficiently credible supporting document found.",
    ),
)


class ReportError(ValueError):
    """The Excel report could not be written."""


@dataclass(frozen=True)
class _CurrencyTotals:
    currency: str
    outgoing: Decimal
    confirmed: Decimal
    probable: Decimal
    missing: Decimal


def write_excel_report(
    result: ReconciliationResult,
    output_path: Path,
    *,
    transactions: Sequence[Transaction] | None = None,
    generated_at: datetime | None = None,
) -> Path:
    """Write an audit workbook for ``result`` and return ``output_path``.

    ``transactions`` is the original statement list. Incoming rows are listed
    from it. When it is omitted, the workbook reports the outgoing rows stored
    on ``result`` and leaves the incoming count empty.

    ``generated_at`` defaults to the current UTC time. Audit rows do not
    depend on it.
    """
    population = _population(result, transactions)
    timestamp = _timestamp(generated_at)
    workbook = Workbook()
    summary = workbook.active
    if summary is None:
        raise ReportError("workbook has no active sheet")
    summary.title = "Summary"
    _build_summary(summary, result, population, transactions is not None, timestamp)
    _build_transactions(workbook.create_sheet("Transactions"), result, population)
    _build_missing(workbook.create_sheet("Missing"), result)
    _build_probable(workbook.create_sheet("Probable"), result)
    _build_unmatched(workbook.create_sheet("Unmatched Docs"), result)
    _build_documents(workbook.create_sheet("Documents"), result)
    duplicates = _reused_documents(result)
    if duplicates:
        _build_duplicates(workbook.create_sheet("Duplicates"), duplicates)
    _save(workbook, Path(output_path))
    return Path(output_path)


def _population(
    result: ReconciliationResult,
    transactions: Sequence[Transaction] | None,
) -> tuple[Transaction, ...]:
    if transactions is None:
        return tuple(row.transaction for row in result.reconciled)
    seen: dict[str, Transaction] = {}
    for transaction in transactions:
        if transaction.id in seen:
            raise ReportError(f"transaction listed twice: {transaction.id}")
        seen[transaction.id] = transaction
    reconciled_ids = {row.transaction.id for row in result.reconciled}
    outgoing_ids = {
        transaction.id
        for transaction in transactions
        if transaction.direction is TransactionDirection.OUTGOING
    }
    if outgoing_ids != reconciled_ids:
        missing = ", ".join(sorted(outgoing_ids - reconciled_ids))
        extra = ", ".join(sorted(reconciled_ids - outgoing_ids))
        raise ReportError(
            "outgoing transactions do not match the reconciliation result: "
            f"missing from result [{missing}]; extra in result [{extra}]"
        )
    for row in result.reconciled:
        supplied = seen[row.transaction.id]
        if supplied != row.transaction:
            raise ReportError(
                f"transaction {row.transaction.id} differs from the reconciliation result"
            )
    return tuple(transactions)


def _timestamp(generated_at: datetime | None) -> datetime:
    moment = generated_at if generated_at is not None else datetime.now(UTC)
    if moment.tzinfo is not None:
        moment = moment.astimezone(UTC).replace(tzinfo=None)
    return moment.replace(microsecond=0)


def _build_summary(
    worksheet: Worksheet,
    result: ReconciliationResult,
    population: Sequence[Transaction],
    incoming_known: bool,
    generated_at: datetime,
) -> None:
    counts = result.summary()
    _check_document_counts(result, counts)
    totals = _currency_totals(result)
    worksheet["A1"] = "Monthly expense audit"
    worksheet["A1"].font = _TITLE_FONT
    row = 3
    row = _section(worksheet, row, "General")
    row = _metric(worksheet, row, "Reporting period", _period(population))
    row = _metric(worksheet, row, "Generated at", generated_at, _TIMESTAMP_FORMAT)
    row = _metric(worksheet, row, "Bank statements", _statements(population))
    row = _metric(worksheet, row, "Total transactions", len(population))
    row = _metric(worksheet, row, "Outgoing transactions", counts.outgoing_count)
    incoming = _incoming_count(population) if incoming_known else None
    row = _metric(worksheet, row, "Incoming transactions", incoming)
    if not incoming_known:
        worksheet.cell(row - 1, 3, "Not included unless the transaction list is supplied")
    unclassified = _direction_count(population, TransactionDirection.UNKNOWN)
    if unclassified:
        row = _metric(worksheet, row, "Unclassified transactions", unclassified)
    row += 1
    row = _section(worksheet, row, "Reconciliation")
    row = _metric(worksheet, row, "Confirmed transactions", counts.confirmed_count)
    row = _metric(worksheet, row, "Probable transactions", counts.probable_count)
    row = _metric(worksheet, row, "Missing transactions", counts.missing_count)
    row = _metric(worksheet, row, "Supporting documents", len(result.document_references))
    matched = len(result.document_references) - counts.unmatched_document_count
    row = _metric(worksheet, row, "Matched documents", matched)
    row = _metric(worksheet, row, "Unmatched documents", counts.unmatched_document_count)
    for item in totals:
        row = _metric(
            worksheet,
            row,
            f"Outgoing amount ({item.currency})",
            item.outgoing,
            _AMOUNT_FORMAT,
        )
        row = _metric(
            worksheet,
            row,
            f"Confirmed amount ({item.currency})",
            item.confirmed,
            _AMOUNT_FORMAT,
        )
        row = _metric(
            worksheet,
            row,
            f"Probable amount ({item.currency})",
            item.probable,
            _AMOUNT_FORMAT,
        )
        row = _metric(
            worksheet,
            row,
            f"Missing amount ({item.currency})",
            item.missing,
            _AMOUNT_FORMAT,
        )
    row += 1
    row = _section(worksheet, row, "Coverage")
    coverage = _ratio(Decimal(counts.confirmed_count), Decimal(counts.outgoing_count))
    row = _metric(worksheet, row, "Coverage by count", coverage, _PERCENT_FORMAT)
    for item in totals:
        amount_coverage = _ratio(item.confirmed, item.outgoing)
        row = _metric(
            worksheet,
            row,
            f"Coverage by amount ({item.currency})",
            amount_coverage,
            _PERCENT_FORMAT,
        )
    row += 1
    row = _section(worksheet, row, "Status")
    for status, meaning in _LEGEND:
        worksheet.cell(row, 1, status)
        worksheet.cell(row, 2, meaning)
        row += 1
    if row > 3:
        _paint_status(worksheet, 1, row - 1)
    _write_count_chart(worksheet, counts)
    _fit_columns(worksheet)
    worksheet.column_dimensions["B"].width = min(
        _MAX_COLUMN_WIDTH,
        max(worksheet.column_dimensions["B"].width or 12, 36),
    )


def _check_document_counts(result: ReconciliationResult, counts: ReconciliationSummary) -> None:
    matched = sum(1 for reference in result.document_references if reference.transactions)
    if matched + counts.unmatched_document_count != len(result.document_references):
        raise ReportError("document usage counts do not match the reconciliation result")


def _currency_totals(result: ReconciliationResult) -> tuple[_CurrencyTotals, ...]:
    outgoing: dict[str, Decimal] = {}
    confirmed: dict[str, Decimal] = {}
    probable: dict[str, Decimal] = {}
    missing: dict[str, Decimal] = {}
    for row in result.reconciled:
        currency = row.transaction.currency
        amount = row.transaction.amount
        outgoing[currency] = outgoing.get(currency, Decimal("0")) + amount
        if row.match.status is MatchStatus.CONFIRMED:
            confirmed[currency] = confirmed.get(currency, Decimal("0")) + amount
        elif row.match.status is MatchStatus.PROBABLE:
            probable[currency] = probable.get(currency, Decimal("0")) + amount
        elif row.match.status is MatchStatus.MISSING:
            missing[currency] = missing.get(currency, Decimal("0")) + amount
        else:
            raise ReportError(f"unknown match status: {row.match.status}")
    totals: list[_CurrencyTotals] = []
    for currency in sorted(outgoing):
        item = _CurrencyTotals(
            currency=currency,
            outgoing=outgoing[currency],
            confirmed=confirmed.get(currency, Decimal("0")),
            probable=probable.get(currency, Decimal("0")),
            missing=missing.get(currency, Decimal("0")),
        )
        parts = item.confirmed + item.probable + item.missing
        if parts != item.outgoing:
            raise ReportError(f"amount totals for {currency} do not add up")
        totals.append(item)
    return tuple(totals)


def _section(worksheet: Worksheet, row: int, title: str) -> int:
    cell = worksheet.cell(row, 1, title)
    cell.font = _SECTION_FONT
    return row + 1


def _metric(
    worksheet: Worksheet,
    row: int,
    label: str,
    value: object,
    number_format: str | None = None,
) -> int:
    worksheet.cell(row, 1, label)
    cell = worksheet.cell(row, 2, value)
    if number_format is not None and value is not None:
        cell.number_format = number_format
    return row + 1


def _period(population: Sequence[Transaction]) -> str | None:
    if not population:
        return None
    first = min(item.posted_date for item in population)
    last = max(item.posted_date for item in population)
    if first == last:
        return first.isoformat()
    return f"{first.isoformat()} to {last.isoformat()}"


def _statements(population: Sequence[Transaction]) -> str | None:
    names = sorted({str(item.source_file) for item in population})
    if not names:
        return None
    return "; ".join(names)


def _incoming_count(population: Sequence[Transaction]) -> int:
    return _direction_count(population, TransactionDirection.INCOMING)


def _direction_count(population: Sequence[Transaction], direction: TransactionDirection) -> int:
    return sum(1 for item in population if item.direction is direction)


def _ratio(numerator: Decimal, denominator: Decimal) -> Decimal | None:
    if denominator == 0:
        return None
    return numerator / denominator


def _write_count_chart(worksheet: Worksheet, counts: ReconciliationSummary) -> None:
    worksheet["E2"] = "Reconciliation counts"
    worksheet["E2"].font = _SECTION_FONT
    worksheet["E3"] = "Status"
    worksheet["F3"] = "Count"
    for cell in (worksheet["E3"], worksheet["F3"]):
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    rows = (
        (MatchStatus.CONFIRMED.value, counts.confirmed_count),
        (MatchStatus.PROBABLE.value, counts.probable_count),
        (MatchStatus.MISSING.value, counts.missing_count),
    )
    for offset, (status, count) in enumerate(rows, start=4):
        worksheet.cell(offset, 5, status)
        worksheet.cell(offset, 6, count)
    chart = BarChart()
    chart.type = "col"
    chart.title = "Reconciliation counts"
    chart.y_axis.title = "Transactions"
    chart.style = 10
    data = Reference(worksheet, min_col=6, min_row=3, max_row=6)
    categories = Reference(worksheet, min_col=5, min_row=4, max_row=6)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)
    chart.shape = 4
    chart.legend = None
    chart.width = 12
    chart.height = 6
    worksheet.add_chart(chart, "H2")


def _build_transactions(
    worksheet: Worksheet,
    result: ReconciliationResult,
    population: Sequence[Transaction],
) -> None:
    _write_header(worksheet, _TRANSACTION_HEADERS)
    rows = {row.transaction.id: row for row in result.reconciled}
    ordered = sorted(population, key=_transaction_key)
    for index, transaction in enumerate(ordered, start=2):
        reconciled = rows.get(transaction.id)
        match = reconciled.match if reconciled is not None else None
        documents = reconciled.documents if reconciled is not None else ()
        _write_table_row(
            worksheet,
            index,
            _TRANSACTION_HEADERS,
            (
                transaction.id,
                transaction.posted_date,
                transaction.direction.value,
                transaction.amount,
                transaction.currency,
                transaction.merchant_raw,
                transaction.merchant_normalized,
                transaction.description,
                match.status.value if match is not None else None,
                match.score if match is not None else None,
                _filenames(documents),
                _document_ids(documents),
                match.reason if match is not None else None,
                match.amount_score if match is not None else None,
                match.currency_score if match is not None else None,
                match.merchant_score if match is not None else None,
                match.date_score if match is not None else None,
                match.reference_score if match is not None else None,
                _combination(match),
                len(documents) if match is not None else None,
            ),
        )
    _finish_table(worksheet, _TRANSACTION_HEADERS, "Status")


def _build_missing(worksheet: Worksheet, result: ReconciliationResult) -> None:
    _write_header(worksheet, _MISSING_HEADERS)
    rows = [row for row in result.reconciled if row.match.status is MatchStatus.MISSING]
    rows.sort(
        key=lambda row: (
            row.transaction.posted_date,
            row.transaction.amount,
            row.transaction.id,
        )
    )
    for index, row in enumerate(rows, start=2):
        transaction = row.transaction
        _write_table_row(
            worksheet,
            index,
            _MISSING_HEADERS,
            (
                transaction.id,
                transaction.posted_date,
                transaction.amount,
                transaction.currency,
                transaction.merchant_raw,
                transaction.description,
                row.match.status.value,
                row.match.score,
                row.match.reason,
            ),
        )
    _finish_table(worksheet, _MISSING_HEADERS, "Status")


def _build_probable(worksheet: Worksheet, result: ReconciliationResult) -> None:
    _write_header(worksheet, _PROBABLE_HEADERS)
    rows = [row for row in result.reconciled if row.match.status is MatchStatus.PROBABLE]
    rows.sort(
        key=lambda row: (
            row.match.score if row.match.score is not None else 2,
            row.transaction.posted_date,
            row.transaction.id,
        )
    )
    for index, row in enumerate(rows, start=2):
        transaction = row.transaction
        _write_table_row(
            worksheet,
            index,
            _PROBABLE_HEADERS,
            (
                transaction.id,
                transaction.posted_date,
                transaction.amount,
                transaction.currency,
                transaction.merchant_raw,
                transaction.description,
                row.match.status.value,
                row.match.score,
                _filenames(row.documents),
                _document_ids(row.documents),
                row.match.reason,
                row.match.amount_score,
                row.match.currency_score,
                row.match.merchant_score,
                row.match.date_score,
                row.match.reference_score,
                _combination(row.match),
                len(row.documents),
            ),
        )
    _finish_table(worksheet, _PROBABLE_HEADERS, "Status")


def _build_unmatched(worksheet: Worksheet, result: ReconciliationResult) -> None:
    _write_header(worksheet, _UNMATCHED_HEADERS)
    documents = sorted(result.unmatched_documents(), key=_document_key)
    for index, document in enumerate(documents, start=2):
        _write_table_row(worksheet, index, _UNMATCHED_HEADERS, _document_values(document))
        _link_filename(worksheet.cell(index, 2), document)
    _finish_table(worksheet, _UNMATCHED_HEADERS, None)


def _build_documents(worksheet: Worksheet, result: ReconciliationResult) -> None:
    _write_header(worksheet, _DOCUMENT_HEADERS)
    references = sorted(result.document_references, key=lambda item: _document_key(item.document))
    for index, reference in enumerate(references, start=2):
        document = reference.document
        usage = _usage(len(reference.transactions))
        transaction_ids = "; ".join(item.id for item in reference.transactions) or None
        _write_table_row(
            worksheet,
            index,
            _DOCUMENT_HEADERS,
            (
                *_document_values(document),
                usage,
                transaction_ids,
            ),
        )
        _link_filename(worksheet.cell(index, 2), document)
    _finish_table(worksheet, _DOCUMENT_HEADERS, None)


def _build_duplicates(
    worksheet: Worksheet,
    references: Sequence[tuple[Document, tuple[str, ...]]],
) -> None:
    _write_header(worksheet, _DUPLICATE_HEADERS)
    ordered = sorted(references, key=lambda item: _document_key(item[0]))
    for index, (document, transaction_ids) in enumerate(ordered, start=2):
        _write_table_row(
            worksheet,
            index,
            _DUPLICATE_HEADERS,
            (
                document.id,
                document.filename,
                document.issued_date,
                document.amount,
                document.currency,
                len(transaction_ids),
                "; ".join(transaction_ids),
            ),
        )
        _link_filename(worksheet.cell(index, 2), document)
    _finish_table(worksheet, _DUPLICATE_HEADERS, None)


def _reused_documents(
    result: ReconciliationResult,
) -> tuple[tuple[Document, tuple[str, ...]], ...]:
    reused: list[tuple[Document, tuple[str, ...]]] = []
    for reference in result.document_references:
        if len(reference.transactions) > 1:
            reused.append((reference.document, tuple(item.id for item in reference.transactions)))
    return tuple(reused)


def _document_values(document: Document) -> tuple[object, ...]:
    method = document.extraction_method.value if document.extraction_method is not None else None
    return (
        document.id,
        document.filename,
        document.document_type.value,
        document.issued_date,
        document.merchant_raw,
        document.merchant_normalized,
        document.amount,
        document.currency,
        document.invoice_number,
        method,
        document.file_hash,
        str(document.filepath),
    )


def _write_header(worksheet: Worksheet, headers: tuple[str, ...]) -> None:
    for column, header in enumerate(headers, start=1):
        cell = worksheet.cell(1, column, header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
    worksheet.row_dimensions[1].height = 18
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}1"


def _write_table_row(
    worksheet: Worksheet,
    row: int,
    headers: tuple[str, ...],
    values: Sequence[object],
) -> None:
    if len(headers) != len(values):
        raise ReportError("report row does not match its headers")
    for column, (header, value) in enumerate(zip(headers, values, strict=True), start=1):
        cell = worksheet.cell(row, column, value)
        if isinstance(value, Decimal) and header in _AMOUNT_HEADERS:
            cell.number_format = _AMOUNT_FORMAT
        elif isinstance(value, datetime):
            cell.number_format = _TIMESTAMP_FORMAT
        elif isinstance(value, date) and header in _DATE_HEADERS:
            cell.number_format = _DATE_FORMAT
        elif (
            isinstance(value, float | int)
            and not isinstance(value, bool)
            and header in _SCORE_HEADERS
        ):
            cell.number_format = _SCORE_FORMAT


def _finish_table(
    worksheet: Worksheet,
    headers: tuple[str, ...],
    status_header: str | None,
) -> None:
    last_row = max(worksheet.max_row, 1)
    worksheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last_row}"
    worksheet.freeze_panes = "A2"
    if status_header is not None and last_row >= 2:
        _paint_status(worksheet, headers.index(status_header) + 1, last_row)
    _fit_columns(worksheet)


def _paint_status(worksheet: Worksheet, column: int, last_row: int) -> None:
    letter = get_column_letter(column)
    span = f"{letter}2:{letter}{last_row}"
    rules = (
        (MatchStatus.CONFIRMED.value, _CONFIRMED_FILL, _CONFIRMED_FONT),
        (MatchStatus.PROBABLE.value, _PROBABLE_FILL, _PROBABLE_FONT),
        (MatchStatus.MISSING.value, _MISSING_FILL, _MISSING_FONT),
    )
    for status, fill, font in rules:
        worksheet.conditional_formatting.add(
            span,
            CellIsRule(operator="equal", formula=[f'"{status}"'], fill=fill, font=font),
        )


def _fit_columns(worksheet: Worksheet) -> None:
    for cells in worksheet.columns:
        letter = get_column_letter(cells[0].column)
        width = 12
        for cell in cells:
            if cell.value is None:
                continue
            width = max(width, min(len(str(cell.value)) + 2, _MAX_COLUMN_WIDTH))
        worksheet.column_dimensions[letter].width = width


def _link_filename(cell: Cell, document: Document) -> None:
    path = document.filepath
    if not path.is_absolute():
        return
    try:
        target = path.as_uri()
    except ValueError:
        return
    cell.hyperlink = target
    cell.style = "Hyperlink"


def _transaction_key(transaction: Transaction) -> tuple[int, date, str]:
    order = {
        TransactionDirection.OUTGOING: 0,
        TransactionDirection.INCOMING: 1,
        TransactionDirection.UNKNOWN: 2,
    }
    return (order[transaction.direction], transaction.posted_date, transaction.id)


def _document_key(document: Document) -> tuple[int, date, str, str]:
    if document.issued_date is None:
        return (1, date.min, document.filename, document.id)
    return (0, document.issued_date, document.filename, document.id)


def _filenames(documents: Sequence[Document]) -> str | None:
    if not documents:
        return None
    return "; ".join(document.filename for document in documents)


def _document_ids(documents: Sequence[Document]) -> str | None:
    if not documents:
        return None
    return "; ".join(document.id for document in documents)


def _combination(match: Match | None) -> str | None:
    if match is None or match.group_amount_exact is None:
        return None
    if match.group_amount_exact:
        return "Yes"
    return "No"


def _usage(count: int) -> str:
    if count <= 0:
        return "UNUSED"
    if count == 1:
        return "USED ONCE"
    return "USED MULTIPLE TIMES"


def _save(workbook: Workbook, output_path: Path) -> None:
    if output_path.suffix.lower() != ".xlsx":
        raise ReportError(f"output path must be an .xlsx file: {output_path}")
    if output_path.exists() and output_path.is_dir():
        raise ReportError(f"output path is a directory: {output_path}")
    parent = output_path.parent
    try:
        parent.mkdir(parents=True, exist_ok=True)
    except OSError as exc:
        raise ReportError(f"cannot create report directory: {parent}") from exc
    temporary = output_path.with_name(f".{output_path.name}.{os.getpid()}.partial")
    try:
        workbook.save(temporary)
        temporary.replace(output_path)
    except Exception as exc:
        temporary.unlink(missing_ok=True)
        raise ReportError(f"cannot write workbook: {output_path}") from exc
