"""Excel report tests use synthetic reconciliation results, not the bank PDF."""

from __future__ import annotations

from datetime import date, datetime, timedelta
from decimal import Decimal
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from expense_auditor.config import load_matching_settings
from expense_auditor.matching.matcher import match_transactions
from expense_auditor.models.document import Document, DocumentType
from expense_auditor.models.match import Match, MatchedDocument, MatchStatus
from expense_auditor.models.reconciliation import (
    ReconciledTransaction,
    ReconciliationResult,
    build_reconciliation_result,
)
from expense_auditor.models.transaction import Transaction, TransactionDirection
from expense_auditor.reporting.excel import ReportError, write_excel_report

_POSTED = date(2026, 9, 10)
_GENERATED = datetime(2026, 10, 5, 12, 0, 0)
_STATEMENT = Path("statement.pdf")


def _transaction(
    transaction_id: str,
    amount: str,
    merchant: str,
    *,
    posted: date = _POSTED,
    currency: str = "EUR",
    direction: TransactionDirection = TransactionDirection.OUTGOING,
    description: str | None = None,
    source_file: Path = _STATEMENT,
) -> Transaction:
    return Transaction.model_validate(
        {
            "id": transaction_id,
            "posted_date": posted,
            "direction": direction,
            "amount": Decimal(amount),
            "currency": currency,
            "description": description or merchant,
            "merchant_raw": merchant,
            "source_file": source_file,
        }
    )


def _document(
    document_id: str,
    amount: str | None,
    merchant: str | None,
    *,
    issued: date | None = _POSTED,
    currency: str | None = "EUR",
    filename: str | None = None,
    invoice_number: str | None = None,
) -> Document:
    name = filename or f"{document_id}.pdf"
    payload: dict[str, object] = {
        "id": document_id,
        "filepath": Path(name),
        "document_type": DocumentType.PDF,
        "issued_date": issued,
        "currency": currency,
        "merchant_raw": merchant,
        "invoice_number": invoice_number,
    }
    if amount is not None:
        payload["amount"] = Decimal(amount)
    return Document.model_validate(payload)


def _confirmed(
    transaction: Transaction,
    documents: tuple[Document, ...],
    *,
    group: bool = False,
    reason: str = "Exact amount and currency; strong merchant match.",
) -> ReconciledTransaction:
    return ReconciledTransaction(
        transaction=transaction,
        match=Match(
            transaction_id=transaction.id,
            status=MatchStatus.CONFIRMED,
            score=load_matching_settings().confirmed_threshold,
            documents=tuple(
                MatchedDocument(
                    document_id=document.id,
                    amount_exact=not group,
                    merchant_similarity=1,
                    date_delta_days=0,
                )
                for document in documents
            ),
            reason=reason,
            amount_score=1,
            currency_score=1,
            merchant_score=1,
            date_score=1,
            reference_score=None,
            group_amount_exact=group,
        ),
        documents=documents,
    )


def _probable(
    transaction: Transaction,
    document: Document,
    score: float,
    reason: str,
) -> ReconciledTransaction:
    return ReconciledTransaction(
        transaction=transaction,
        match=Match(
            transaction_id=transaction.id,
            status=MatchStatus.PROBABLE,
            score=score,
            documents=(
                MatchedDocument(
                    document_id=document.id,
                    amount_exact=True,
                    merchant_similarity=0.75,
                    date_delta_days=5,
                ),
            ),
            reason=reason,
            amount_score=1,
            currency_score=1,
            merchant_score=0.75,
            date_score=0.2,
            reference_score=None,
            group_amount_exact=False,
        ),
        documents=(document,),
    )


def _missing(transaction: Transaction, reason: str) -> ReconciledTransaction:
    return ReconciledTransaction(
        transaction=transaction,
        match=Match(
            transaction_id=transaction.id,
            status=MatchStatus.MISSING,
            score=0,
            reason=reason,
        ),
        documents=(),
    )


def _summary_value(worksheet: Worksheet, label: str) -> object:
    for row in worksheet.iter_rows(min_col=1, max_col=2):
        if row[0].value == label:
            return row[1].value
    raise AssertionError(f"missing summary label: {label}")


def _headers(worksheet: Worksheet) -> dict[str, int]:
    return {str(cell.value): cell.column for cell in worksheet[1] if cell.value is not None}


def _rows(worksheet: Worksheet) -> list[dict[str, object]]:
    headers = _headers(worksheet)
    found: list[dict[str, object]] = []
    for row_number in range(2, worksheet.max_row + 1):
        values = {
            header: worksheet.cell(row_number, column).value
            for header, column in headers.items()
        }
        if any(value is not None for value in values.values()):
            found.append(values)
    return found


def _money(value: object) -> Decimal:
    assert isinstance(value, int | float)
    return Decimal(str(value)).quantize(Decimal("0.01"))


def _day(value: object) -> date:
    assert isinstance(value, datetime)
    return value.date()


def _sample() -> tuple[ReconciliationResult, tuple[Transaction, ...]]:
    confirmed = _transaction("txn-confirmed", "40.00", "Amazon", posted=date(2026, 9, 2))
    probable = _transaction("txn-probable", "25.00", "ALPHA SUPPLIES", posted=date(2026, 9, 8))
    missing_late = _transaction("txn-missing-late", "15.00", "GROCERY", posted=date(2026, 9, 12))
    missing_early = _transaction("txn-missing-early", "8.00", "CAFE", posted=date(2026, 9, 3))
    incoming = _transaction(
        "txn-incoming",
        "500.00",
        "Client",
        posted=date(2026, 9, 1),
        direction=TransactionDirection.INCOMING,
    )
    confirmed_doc = _document("doc-amazon", "40.00", "Amazon", filename="amazon.pdf")
    probable_doc = _document("doc-beta", "25.00", "BETA SUPPLIES", filename="beta.pdf")
    unused = _document(
        "doc-orphan",
        "9.99",
        "Unused Shop",
        filename="orphan.pdf",
        issued=date(2026, 9, 20),
    )
    rows = (
        _confirmed(confirmed, (confirmed_doc,)),
        _probable(probable, probable_doc, 0.72, "Merchant similarity is moderate."),
        _missing(missing_late, "No supporting document found with sufficient evidence."),
        _missing(missing_early, "No supporting document found with sufficient evidence."),
    )
    documents = (unused, probable_doc, confirmed_doc)
    result = build_reconciliation_result(rows, documents)
    return result, (incoming, missing_late, confirmed, probable, missing_early)


def test_summary_counts_amounts_and_coverage(tmp_path: Path) -> None:
    result, transactions = _sample()
    path = write_excel_report(
        result,
        tmp_path / "audit.xlsx",
        transactions=transactions,
        generated_at=_GENERATED,
    )
    summary = load_workbook(path)["Summary"]

    assert _summary_value(summary, "Reporting period") == "2026-09-01 to 2026-09-12"
    assert _summary_value(summary, "Generated at") == _GENERATED
    assert _summary_value(summary, "Bank statements") == "statement.pdf"
    assert _summary_value(summary, "Total transactions") == 5
    assert _summary_value(summary, "Outgoing transactions") == 4
    assert _summary_value(summary, "Incoming transactions") == 1
    assert _summary_value(summary, "Confirmed transactions") == 1
    assert _summary_value(summary, "Probable transactions") == 1
    assert _summary_value(summary, "Missing transactions") == 2
    assert _summary_value(summary, "Supporting documents") == 3
    assert _summary_value(summary, "Matched documents") == 2
    assert _summary_value(summary, "Unmatched documents") == 1
    assert _money(_summary_value(summary, "Outgoing amount (EUR)")) == Decimal("88.00")
    assert _money(_summary_value(summary, "Confirmed amount (EUR)")) == Decimal("40.00")
    assert _money(_summary_value(summary, "Probable amount (EUR)")) == Decimal("25.00")
    assert _money(_summary_value(summary, "Missing amount (EUR)")) == Decimal("23.00")
    assert _summary_value(summary, "Coverage by count") == 0.25
    coverage = _summary_value(summary, "Coverage by amount (EUR)")
    assert isinstance(coverage, float)
    assert Decimal(str(coverage)).quantize(Decimal("0.0001")) == Decimal("0.4545")
    coverage_cell = next(
        row[1]
        for row in summary.iter_rows(min_col=1, max_col=2)
        if row[0].value == "Coverage by count"
    )
    assert coverage_cell.number_format == "0.0%"
    assert summary["E4"].value == "CONFIRMED"
    assert summary["F4"].value == 1
    assert summary["F5"].value == 1
    assert summary["F6"].value == 2
    assert len(summary._charts) == 1
    assert "CONFIRMED" in " ".join(str(cell.value) for row in summary.iter_rows() for cell in row)
    assert "Plausible match requiring review" in " ".join(
        str(cell.value) for row in summary.iter_rows() for cell in row if cell.value
    )


def test_transactions_sheet_lists_rows_and_evidence(tmp_path: Path) -> None:
    result, transactions = _sample()
    path = write_excel_report(result, tmp_path / "audit.xlsx", transactions=transactions)
    sheet = load_workbook(path)["Transactions"]
    rows = _rows(sheet)

    assert [row["Transaction ID"] for row in rows] == [
        "txn-confirmed",
        "txn-missing-early",
        "txn-probable",
        "txn-missing-late",
        "txn-incoming",
    ]
    confirmed = rows[0]
    assert confirmed["Status"] == "CONFIRMED"
    assert confirmed["Matched Documents"] == "amazon.pdf"
    assert confirmed["Matched Document IDs"] == "doc-amazon"
    assert confirmed["Amount Score"] == 1
    assert confirmed["Merchant Score"] == 1
    assert confirmed["Date Score"] == 1
    assert confirmed["Currency Score"] == 1
    assert confirmed["Combination Match"] == "No"
    assert confirmed["Number of Documents"] == 1
    assert _money(confirmed["Amount"]) == Decimal("40.00")
    assert _day(confirmed["Date"]) == date(2026, 9, 2)
    assert rows[-1]["Transaction ID"] == "txn-incoming"
    assert rows[-1]["Status"] is None
    assert rows[-1]["Direction"] == "incoming"
    formulas: list[str] = []
    for formatted in sheet.conditional_formatting:
        for rule in sheet.conditional_formatting[formatted]:
            formulas.extend(rule.formula)
    assert '"CONFIRMED"' in formulas
    assert '"PROBABLE"' in formulas
    assert '"MISSING"' in formulas
    amount_cell = sheet.cell(2, _headers(sheet)["Amount"])
    date_cell = sheet.cell(2, _headers(sheet)["Date"])
    assert amount_cell.data_type == "n"
    assert amount_cell.number_format == "#,##0.00"
    assert date_cell.is_date
    assert date_cell.number_format == "YYYY-MM-DD"
    assert sheet.freeze_panes == "A2"
    assert sheet.auto_filter.ref is not None
    assert sheet.auto_filter.ref.startswith("A1:")
    assert sheet["A1"].font.bold is True
    assert all((dimension.width or 0) <= 48 for dimension in sheet.column_dimensions.values())


def test_missing_sheet_lists_only_missing_and_sorts_by_date_then_amount(tmp_path: Path) -> None:
    early_small = _transaction("txn-early-small", "5.00", "CAFE", posted=date(2026, 9, 3))
    early_large = _transaction("txn-early-large", "12.00", "CAFE", posted=date(2026, 9, 3))
    later = _transaction("txn-later", "1.00", "CAFE", posted=date(2026, 9, 4))
    confirmed = _transaction("txn-confirmed", "40.00", "Amazon")
    document = _document("doc-amazon", "40.00", "Amazon")
    result = build_reconciliation_result(
        (
            _missing(later, "later"),
            _confirmed(confirmed, (document,)),
            _missing(early_large, "early large"),
            _missing(early_small, "early small"),
        ),
        (document,),
    )
    path = write_excel_report(result, tmp_path / "audit.xlsx", generated_at=_GENERATED)
    rows = _rows(load_workbook(path)["Missing"])

    assert [row["Transaction ID"] for row in rows] == [
        "txn-early-small",
        "txn-early-large",
        "txn-later",
    ]
    assert {row["Status"] for row in rows} == {"MISSING"}
    assert rows[0]["Reason"] == "early small"


def test_probable_sheet_sorts_weakest_first(tmp_path: Path) -> None:
    weak = _transaction("txn-weak", "10.00", "ALPHA SUPPLIES", posted=date(2026, 9, 2))
    strong = _transaction("txn-strong", "11.00", "ALPHA SUPPLIES", posted=date(2026, 9, 1))
    weak_doc = _document("doc-weak", "10.00", "BETA SUPPLIES")
    strong_doc = _document("doc-strong", "11.00", "BETA SUPPLIES")
    confirmed = _transaction("txn-confirmed", "40.00", "Amazon")
    confirmed_doc = _document("doc-amazon", "40.00", "Amazon")
    result = build_reconciliation_result(
        (
            _probable(strong, strong_doc, 0.8, "stronger probable"),
            _confirmed(confirmed, (confirmed_doc,)),
            _probable(weak, weak_doc, 0.61, "weaker probable"),
        ),
        (weak_doc, strong_doc, confirmed_doc),
    )
    path = write_excel_report(result, tmp_path / "audit.xlsx", generated_at=_GENERATED)
    rows = _rows(load_workbook(path)["Probable"])

    assert [row["Transaction ID"] for row in rows] == ["txn-weak", "txn-strong"]
    assert rows[0]["Match Score"] == 0.61
    assert rows[0]["Matched Documents"] == "doc-weak.pdf"
    assert rows[0]["Merchant Score"] == 0.75
    assert {row["Status"] for row in rows} == {"PROBABLE"}


def test_unmatched_docs_exclude_used_documents(tmp_path: Path) -> None:
    result, transactions = _sample()
    path = write_excel_report(
        result,
        tmp_path / "audit.xlsx",
        transactions=transactions,
        generated_at=_GENERATED,
    )
    rows = _rows(load_workbook(path)["Unmatched Docs"])

    assert [row["Document ID"] for row in rows] == ["doc-orphan"]
    assert rows[0]["Filename"] == "orphan.pdf"
    assert _money(rows[0]["Amount"]) == Decimal("9.99")
    assert rows[0]["Currency"] == "EUR"
    assert rows[0]["Invoice Number"] is None


def test_documents_inventory_records_usage(tmp_path: Path) -> None:
    result, transactions = _sample()
    path = write_excel_report(
        result,
        tmp_path / "audit.xlsx",
        transactions=transactions,
        generated_at=_GENERATED,
    )
    rows = {row["Document ID"]: row for row in _rows(load_workbook(path)["Documents"])}

    assert set(rows) == {"doc-amazon", "doc-beta", "doc-orphan"}
    assert rows["doc-amazon"]["Match Usage"] == "USED ONCE"
    assert rows["doc-amazon"]["Matched Transaction IDs"] == "txn-confirmed"
    assert rows["doc-amazon"]["Merchant Raw"] == "Amazon"
    assert rows["doc-orphan"]["Match Usage"] == "UNUSED"
    assert rows["doc-orphan"]["Matched Transaction IDs"] is None
    assert [row["Document ID"] for row in _rows(load_workbook(path)["Documents"])] == [
        "doc-amazon",
        "doc-beta",
        "doc-orphan",
    ]


def test_multiple_documents_stay_on_one_row(tmp_path: Path) -> None:
    transaction = _transaction("txn-combo", "100.00", "Office Depot")
    first = _document("doc-a", "40.00", "Office Depot", filename="receipt_a.pdf")
    second = _document("doc-b", "60.00", "Office Depot", filename="receipt_b.pdf")
    result = build_reconciliation_result(
        (_confirmed(transaction, (first, second), group=True),),
        (second, first),
    )
    path = write_excel_report(result, tmp_path / "audit.xlsx", generated_at=_GENERATED)
    workbook = load_workbook(path)
    rows = _rows(workbook["Transactions"])

    assert len(rows) == 1
    assert rows[0]["Matched Documents"] == "receipt_a.pdf; receipt_b.pdf"
    assert rows[0]["Matched Document IDs"] == "doc-a; doc-b"
    assert rows[0]["Combination Match"] == "Yes"
    assert rows[0]["Number of Documents"] == 2
    assert "Duplicates" not in workbook.sheetnames


def test_currencies_stay_separate(tmp_path: Path) -> None:
    euro = _transaction("txn-eur", "10.50", "Amazon", currency="EUR")
    dollar = _transaction("txn-usd", "3.25", "Amazon", currency="USD")
    euro_doc = _document("doc-eur", "10.50", "Amazon", currency="EUR")
    dollar_doc = _document("doc-usd", "3.25", "Amazon", currency="USD")
    result = build_reconciliation_result(
        (
            _confirmed(euro, (euro_doc,)),
            _probable(dollar, dollar_doc, 0.7, "dollar probable"),
        ),
        (euro_doc, dollar_doc),
    )
    path = write_excel_report(result, tmp_path / "audit.xlsx", generated_at=_GENERATED)
    workbook = load_workbook(path)
    summary = workbook["Summary"]
    transactions = _rows(workbook["Transactions"])

    assert _money(_summary_value(summary, "Outgoing amount (EUR)")) == Decimal("10.50")
    assert _money(_summary_value(summary, "Confirmed amount (EUR)")) == Decimal("10.50")
    assert _money(_summary_value(summary, "Outgoing amount (USD)")) == Decimal("3.25")
    assert _money(_summary_value(summary, "Probable amount (USD)")) == Decimal("3.25")
    assert _money(_summary_value(summary, "Confirmed amount (USD)")) == Decimal("0.00")
    amounts = {row["Currency"]: _money(row["Amount"]) for row in transactions}
    assert amounts == {"EUR": Decimal("10.50"), "USD": Decimal("3.25")}


def test_empty_result_is_a_valid_workbook(tmp_path: Path) -> None:
    result = ReconciliationResult(reconciled=(), document_references=())
    path = write_excel_report(
        result,
        tmp_path / "empty.xlsx",
        transactions=[],
        generated_at=_GENERATED,
    )
    workbook = load_workbook(path)

    assert workbook.sheetnames == [
        "Summary",
        "Transactions",
        "Missing",
        "Probable",
        "Unmatched Docs",
        "Documents",
    ]
    summary = workbook["Summary"]
    assert _summary_value(summary, "Total transactions") == 0
    assert _summary_value(summary, "Outgoing transactions") == 0
    assert _summary_value(summary, "Incoming transactions") == 0
    assert _summary_value(summary, "Coverage by count") is None
    assert _summary_value(summary, "Reporting period") is None
    assert summary["F4"].value == 0
    for name in ("Transactions", "Missing", "Probable", "Unmatched Docs", "Documents"):
        sheet = workbook[name]
        assert sheet.max_row == 1
        assert sheet.freeze_panes == "A2"
        assert _rows(sheet) == []


def test_omitted_transaction_list_leaves_incoming_blank(tmp_path: Path) -> None:
    transaction = _transaction("txn-out", "10.00", "Amazon")
    document = _document("doc-amazon", "10.00", "Amazon")
    result = build_reconciliation_result((_confirmed(transaction, (document,)),), (document,))
    path = write_excel_report(result, tmp_path / "audit.xlsx", generated_at=_GENERATED)
    summary = load_workbook(path)["Summary"]

    assert _summary_value(summary, "Incoming transactions") is None
    assert _summary_value(summary, "Total transactions") == 1
    assert _rows(load_workbook(path)["Missing"]) == []
    assert _rows(load_workbook(path)["Probable"]) == []
    assert _rows(load_workbook(path)["Unmatched Docs"]) == []


def test_duplicate_sheet_appears_when_a_document_is_reused(tmp_path: Path) -> None:
    shared = _document("doc-shared", "10.00", "Amazon", filename="shared.pdf")
    first = _transaction("txn-a", "10.00", "Amazon", posted=date(2026, 9, 1))
    second = _transaction("txn-b", "10.00", "Amazon", posted=date(2026, 9, 2))
    result = build_reconciliation_result(
        (
            _confirmed(first, (shared,)),
            _probable(second, shared, 0.7, "same receipt"),
        ),
        (shared,),
    )
    path = write_excel_report(result, tmp_path / "audit.xlsx", generated_at=_GENERATED)
    workbook = load_workbook(path)
    documents = _rows(workbook["Documents"])
    duplicates = _rows(workbook["Duplicates"])

    assert documents[0]["Match Usage"] == "USED MULTIPLE TIMES"
    assert documents[0]["Matched Transaction IDs"] == "txn-a; txn-b"
    assert duplicates[0]["Document ID"] == "doc-shared"
    assert duplicates[0]["Transaction Count"] == 2


def test_absolute_path_is_linked_and_relative_path_stays_text(tmp_path: Path) -> None:
    absolute = tmp_path / "receipt-a.pdf"
    absolute.write_bytes(b"receipt")
    linked = Document.model_validate(
        {
            "id": "doc-abs",
            "filepath": absolute,
            "document_type": DocumentType.PDF,
            "merchant_raw": "Amazon",
            "amount": Decimal("10.00"),
            "currency": "EUR",
            "issued_date": _POSTED,
        }
    )
    relative = _document("doc-rel", "9.00", "Unused Shop", filename="relative.pdf")
    transaction = _transaction("txn-out", "10.00", "Amazon")
    result = build_reconciliation_result((_confirmed(transaction, (linked,)),), (linked, relative))
    path = write_excel_report(result, tmp_path / "audit.xlsx", generated_at=_GENERATED)
    sheet = load_workbook(path)["Documents"]
    rows = _rows(sheet)
    absolute_row = next(
        index for index, row in enumerate(rows, start=2) if row["Document ID"] == "doc-abs"
    )
    relative_row = next(
        index for index, row in enumerate(rows, start=2) if row["Document ID"] == "doc-rel"
    )
    filename_column = _headers(sheet)["Filename"]

    assert sheet.cell(absolute_row, filename_column).hyperlink is not None
    assert sheet.cell(absolute_row, filename_column).hyperlink.target == absolute.as_uri()
    assert sheet.cell(absolute_row, filename_column).value == absolute.name
    assert sheet.cell(relative_row, filename_column).hyperlink is None
    assert sheet.cell(relative_row, filename_column).value == "relative.pdf"


def test_same_result_writes_the_same_audit_rows(tmp_path: Path) -> None:
    result, transactions = _sample()
    first = write_excel_report(
        result,
        tmp_path / "first.xlsx",
        transactions=transactions,
        generated_at=_GENERATED,
    )
    second = write_excel_report(
        result,
        tmp_path / "second.xlsx",
        transactions=transactions,
        generated_at=_GENERATED,
    )
    left = load_workbook(first)
    right = load_workbook(second)

    assert left.sheetnames == right.sheetnames
    for name in left.sheetnames:
        left_rows = list(left[name].iter_rows(values_only=True))
        right_rows = list(right[name].iter_rows(values_only=True))
        assert left_rows == right_rows


def test_invalid_output_and_inconsistent_data_fail(tmp_path: Path) -> None:
    transaction = _transaction("txn-out", "10.00", "Amazon")
    document = _document("doc-amazon", "10.00", "Amazon")
    result = build_reconciliation_result((_confirmed(transaction, (document,)),), (document,))
    existing = tmp_path / "audit.xlsx"
    write_excel_report(result, existing, generated_at=_GENERATED)
    original = existing.read_bytes()
    blocked = tmp_path / "not-a-directory"
    blocked.write_text("keep", encoding="utf-8")

    with pytest.raises(ReportError):
        write_excel_report(result, blocked / "audit.xlsx", generated_at=_GENERATED)
    with pytest.raises(ReportError):
        write_excel_report(result, tmp_path / "audit.csv", generated_at=_GENERATED)
    extra = _transaction("txn-extra", "4.00", "Other")
    with pytest.raises(ReportError, match="outgoing transactions"):
        write_excel_report(result, tmp_path / "other.xlsx", transactions=(transaction, extra))

    assert blocked.read_text(encoding="utf-8") == "keep"
    assert existing.read_bytes() == original
    assert not (tmp_path / "other.xlsx").exists()


def test_failed_save_keeps_the_existing_workbook(
    tmp_path: Path,
    monkeypatch: pytest.MonkeyPatch,
) -> None:
    transaction = _transaction("txn-out", "10.00", "Amazon")
    document = _document("doc-amazon", "10.00", "Amazon")
    result = build_reconciliation_result((_confirmed(transaction, (document,)),), (document,))
    path = tmp_path / "audit.xlsx"
    write_excel_report(result, path, generated_at=_GENERATED)
    original = path.read_bytes()

    def fail_save(self: Workbook, filename: str) -> None:
        raise OSError("disk full")

    monkeypatch.setattr(Workbook, "save", fail_save)
    with pytest.raises(ReportError, match="cannot write workbook"):
        write_excel_report(result, path, generated_at=_GENERATED)

    assert path.read_bytes() == original
    assert list(tmp_path.glob("*.partial")) == []


def test_monthly_reconciliation_report(tmp_path: Path) -> None:
    posted = date(2026, 9, 10)
    confirmed = _transaction("txn-confirmed", "103.51", "AMAZON* NV3LD00P4", posted=posted)
    probable = _transaction("txn-probable", "42.50", "ALPHA SUPPLIES", posted=posted)
    missing = _transaction("txn-missing", "77.10", "GROCERY STORE", posted=posted)
    combo = _transaction("txn-combo", "100.00", "Office Depot", posted=posted)
    incoming = _transaction(
        "txn-incoming",
        "1834.71",
        "Client Payment",
        direction=TransactionDirection.INCOMING,
        posted=posted - timedelta(days=2),
    )
    documents = (
        _document("doc-amazon", "103.51", "Amazon", filename="amazon.pdf", issued=posted),
        _document(
            "doc-beta",
            "42.50",
            "BETA SUPPLIES",
            filename="beta.pdf",
            issued=posted - timedelta(days=5),
        ),
        _document("doc-20", "20.00", "Office Depot", filename="part-20.pdf", issued=posted),
        _document("doc-30", "30.00", "Office Depot", filename="part-30.pdf", issued=posted),
        _document("doc-50", "50.00", "Office Depot", filename="part-50.pdf", issued=posted),
        _document("doc-orphan", "9.99", "Unused Shop", filename="orphan.pdf", issued=posted),
    )
    transactions = (confirmed, probable, missing, combo, incoming)
    result = match_transactions(transactions, documents)
    path = write_excel_report(
        result,
        tmp_path / "september.xlsx",
        transactions=transactions,
        generated_at=_GENERATED,
    )
    workbook = load_workbook(path)
    summary = workbook["Summary"]
    transaction_rows = {row["Transaction ID"]: row for row in _rows(workbook["Transactions"])}

    assert _summary_value(summary, "Outgoing transactions") == 4
    assert _summary_value(summary, "Incoming transactions") == 1
    assert _summary_value(summary, "Confirmed transactions") == 2
    assert _summary_value(summary, "Probable transactions") == 1
    assert _summary_value(summary, "Missing transactions") == 1
    assert _summary_value(summary, "Unmatched documents") == 1
    assert transaction_rows["txn-confirmed"]["Status"] == "CONFIRMED"
    assert transaction_rows["txn-probable"]["Status"] == "PROBABLE"
    assert transaction_rows["txn-missing"]["Status"] == "MISSING"
    assert transaction_rows["txn-incoming"]["Status"] is None
    assert transaction_rows["txn-combo"]["Combination Match"] == "Yes"
    combo_names = str(transaction_rows["txn-combo"]["Matched Documents"])
    combo_ids = str(transaction_rows["txn-combo"]["Matched Document IDs"])
    assert set(combo_names.split("; ")) == {"part-20.pdf", "part-30.pdf", "part-50.pdf"}
    assert set(combo_ids.split("; ")) == {"doc-20", "doc-30", "doc-50"}
    assert [row["Transaction ID"] for row in _rows(workbook["Missing"])] == ["txn-missing"]
    assert [row["Transaction ID"] for row in _rows(workbook["Probable"])] == ["txn-probable"]
    assert [row["Document ID"] for row in _rows(workbook["Unmatched Docs"])] == ["doc-orphan"]
    assert len(_rows(workbook["Documents"])) == 6
